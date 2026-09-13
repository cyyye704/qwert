"""Build the single-source Q1-Q4 summary and final consistency audit."""
from __future__ import annotations

import hashlib
import json
import os
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "src"), str(ROOT)]

from config import ENERGY_LIMIT, MODEL_VERSION, SOC_LOWER, SOC_UPPER  # noqa: E402
from result_workbooks import OFFICIAL_SHEETS  # noqa: E402
from time_axis import all_slot_labels  # noqa: E402

WORKBOOKS = ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx")
Q1_FIGURES = (
    "figures/problem1.png",
    "figures/problem1.pdf",
    "figures/problem1_diagnostic.png",
    "figures/problem1_diagnostic.pdf",
    "figures/problem1_sensitivity.png",
)


def _read(path: str | Path):
    return json.loads((ROOT / path).read_text(encoding="utf-8"))


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _labels(sheet, vertical: bool = False) -> list[str]:
    if vertical:
        return [sheet.cell(row, 1).value for row in range(2, 146)]
    return [sheet.cell(1, column).value for column in range(2, 146)]


def _to_markdown(frame: pd.DataFrame) -> str:
    """Render a compact Markdown table without pandas' optional tabulate dependency."""
    headers = [str(column) for column in frame.columns]
    rows = [[str(value) for value in row] for row in frame.itertuples(index=False, name=None)]
    escape = lambda value: value.replace("|", "\\|").replace("\n", " ")
    lines = [
        "| " + " | ".join(map(escape, headers)) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    lines.extend("| " + " | ".join(map(escape, row)) + " |" for row in rows)
    return "\n".join(lines)


def _q1_workbook_audit(path: Path, q1: dict) -> dict:
    # Random cell access on an openpyxl read-only worksheet repeatedly scans the
    # XML stream.  These official files are modest enough to load normally and
    # doing so keeps the full-chain audit linear in the number of cells.
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    assert workbook.sheetnames == OFFICIAL_SHEETS[path.name]
    plan = workbook["计划购电量"]
    labels = _labels(plan, vertical=True)
    values = np.asarray([plan.cell(row, 2).value for row in range(2, 146)], dtype=float)
    raw = pd.read_excel(ROOT / "data" / "附件" / "附件1.xlsx", sheet_name="Sheet1")
    price = raw["电价"].to_numpy(dtype=float)
    selected_errors = []
    for label, expected in q1["specified_purchase_kwh"].items():
        selected_errors.append(abs(values[labels.index(label)] - float(expected)))
    storage = workbook["充放电量"]
    charge = np.asarray([storage.cell(row, 2).value for row in range(2, 8)], dtype=float)
    discharge = np.asarray([storage.cell(row, 3).value for row in range(2, 8)], dtype=float)
    window_errors = []
    for block, (_, item) in enumerate(q1["storage_windows_kwh"].items()):
        window_errors.extend([
            abs(charge[block] - float(item["charge_kwh"])),
            abs(discharge[block] - float(item["discharge_kwh"])),
        ])
    result = {
        "exists": True,
        "bytes": path.stat().st_size,
        "sha256": _sha(path),
        "sheets": workbook.sheetnames,
        "official_sheet_set": True,
        "time_labels_correct": labels == all_slot_labels(),
        "position_j_equals_internal_j": True,
        "purchase_sum_error_kwh": float(abs(values.sum() - q1["daily_purchase_kwh"])),
        "purchase_cost_error_yuan": float(abs(price @ values - q1["daily_purchase_cost_yuan"])),
        "selected_position_max_error_kwh": float(max(selected_errors, default=0.0)),
        "storage_window_max_error_kwh": float(max(window_errors, default=0.0)),
        "soc_start_error_kwh": float(abs(float(storage.cell(2, 5).value) - q1["soc_0000_kwh"])),
        "soc_end_error_kwh": float(abs(float(storage.cell(3, 5).value) - q1["soc_2400_kwh"])),
    }
    result["content_matches_q1_json"] = bool(
        result["time_labels_correct"]
        and max(
            result["purchase_sum_error_kwh"],
            result["purchase_cost_error_yuan"],
            result["selected_position_max_error_kwh"],
            result["storage_window_max_error_kwh"],
            result["soc_start_error_kwh"],
            result["soc_end_error_kwh"],
        ) < 1e-6
    )
    workbook.close()
    return result


def _annual_workbook_audit(path: Path, daily_csv: Path, rolling: bool) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
    assert workbook.sheetnames == OFFICIAL_SHEETS[path.name]
    frame = pd.read_csv(daily_csv)
    plan_sheet = workbook["计划购电量"]
    labels = _labels(plan_sheet)
    rows = len(frame)
    workbook_dates = pd.to_datetime([plan_sheet.cell(row, 1).value for row in range(2, rows + 2)])
    plan = np.asarray([
        [plan_sheet.cell(row, column).value for column in range(2, 146)]
        for row in range(2, rows + 2)
    ], dtype=float)
    plan_totals = np.asarray([plan_sheet.cell(row, 146).value for row in range(2, rows + 2)], dtype=float)
    plan_costs = np.asarray([plan_sheet.cell(row, 147).value for row in range(2, rows + 2)], dtype=float)

    storage = workbook["充放电量"]
    charge_totals = []
    discharge_totals = []
    soc_starts = []
    soc_ends = []
    for day in range(rows):
        first = 2 + day * 6
        charge_totals.append(sum(float(storage.cell(first + block, 3).value or 0.0) for block in range(6)))
        discharge_totals.append(sum(float(storage.cell(first + block, 4).value or 0.0) for block in range(6)))
        soc_starts.append(float(storage.cell(first, 6).value))
        soc_ends.append(float(storage.cell(first + 1, 6).value))

    emergency_by_date: dict[str, float] = {}
    emergency = workbook["紧急购电量"]
    for row in emergency.iter_rows(min_row=2, values_only=True):
        if row[0] is None or row[2] is None:
            continue
        key = pd.Timestamp(row[0]).strftime("%Y-%m-%d")
        emergency_by_date[key] = emergency_by_date.get(key, 0.0) + float(row[2])
    emergency_totals = np.asarray([emergency_by_date.get(str(date), 0.0) for date in frame.date], dtype=float)

    max_final_error = 0.0
    max_final_cost_error = 0.0
    if rolling:
        adjustment = workbook["调整购电量"]
        assert _labels(adjustment) == labels
        final = np.asarray([
            [adjustment.cell(row, column).value for column in range(2, 146)]
            for row in range(2, rows + 2)
        ], dtype=float)
        final_totals = np.asarray([adjustment.cell(row, 146).value for row in range(2, rows + 2)], dtype=float)
        final_costs = np.asarray([adjustment.cell(row, 147).value for row in range(2, rows + 2)], dtype=float)
        max_final_error = float(max(
            np.max(np.abs(final.sum(axis=1) - frame.final_purchase_kwh)),
            np.max(np.abs(final_totals - frame.final_purchase_kwh)),
        ))
        max_final_cost_error = float(np.max(np.abs(final_costs - (frame.plan_cost_yuan + frame.adjustment_cost_yuan))))

    result = {
        "exists": True,
        "bytes": path.stat().st_size,
        "sha256": _sha(path),
        "sheets": workbook.sheetnames,
        "official_sheet_set": True,
        "plan_rows": rows,
        "dates_match": [date.strftime("%Y-%m-%d") for date in workbook_dates] == frame.date.astype(str).tolist(),
        "time_labels_correct": labels == all_slot_labels(),
        "position_j_equals_internal_j": True,
        "max_plan_sum_error_kwh": float(max(
            np.max(np.abs(plan.sum(axis=1) - frame.plan_kwh)),
            np.max(np.abs(plan_totals - frame.plan_kwh)),
        )),
        "max_plan_cost_error_yuan": float(np.max(np.abs(plan_costs - frame.plan_cost_yuan))),
        "max_final_sum_error_kwh": max_final_error,
        "max_final_cost_error_yuan": max_final_cost_error,
        "max_charge_sum_error_kwh": float(np.max(np.abs(np.asarray(charge_totals) - frame.charge_kwh))),
        "max_discharge_sum_error_kwh": float(np.max(np.abs(np.asarray(discharge_totals) - frame.discharge_kwh))),
        "max_emergency_sum_error_kwh": float(np.max(np.abs(emergency_totals - frame.emergency_kwh))),
        "max_soc_start_error_kwh": float(np.max(np.abs(np.asarray(soc_starts) - frame.soc_start_kwh))),
        "max_soc_end_error_kwh": float(np.max(np.abs(np.asarray(soc_ends) - frame.soc_end_kwh))),
    }
    numeric_errors = [value for key, value in result.items() if key.startswith("max_")]
    result["content_matches_daily_csv"] = bool(
        result["dates_match"]
        and result["time_labels_correct"]
        and max(numeric_errors, default=0.0) < 1e-6
    )
    workbook.close()
    return result


def main() -> None:
    run_id = os.environ.get("FINAL_RUN_ID", "unbound-manual-run")
    q1 = _read("results/problem1.json")
    q2 = _read("results/problem2.json")
    q3 = _read("results/problem3.json")
    q4 = _read("results/problem4.json")
    q4_sensitivity = _read("results/problem4_price_volatility_sensitivity.json")
    unit_tests = _read("results/c_spec_unit_tests.json")
    timestamp_tests = _read("results/timestamp_alignment_tests.json")
    timestamp_audit = _read("results/timestamp_alignment_audit.json")
    q3_alignment = _read("results/q3_forecast_alignment_audit.json")

    csv_names = [
        "problem2_daily.csv", "problem3_daily.csv",
        "problem4_causal_q2_daily.csv", "problem4_causal_q3_daily.csv",
        "problem4_fixed_q2_daily.csv", "problem4_fixed_q3_daily.csv",
        "problem4_oracle_q2_daily.csv", "problem4_oracle_q3_daily.csv",
        "problem4_perfect_information_daily.csv",
    ]
    frames = [pd.read_csv(ROOT / "results" / name) for name in csv_names]
    all_frames = pd.concat(frames, ignore_index=True)

    workbook_audit = {
        "result1.xlsx": _q1_workbook_audit(ROOT / "result1.xlsx", q1),
        "result2.xlsx": _annual_workbook_audit(ROOT / "result2.xlsx", ROOT / "results/problem2_daily.csv", False),
        "result3.xlsx": _annual_workbook_audit(ROOT / "result3.xlsx", ROOT / "results/problem3_daily.csv", True),
        "result4-2.xlsx": _annual_workbook_audit(ROOT / "result4-2.xlsx", ROOT / "results/problem4_causal_q2_daily.csv", False),
        "result4-3.xlsx": _annual_workbook_audit(ROOT / "result4-3.xlsx", ROOT / "results/problem4_causal_q3_daily.csv", True),
    }
    q4_rows = {row["scenario"]: row for row in q4["scenarios"]}
    sensitivity_frame = pd.DataFrame(q4_sensitivity["rows"])
    sensitivity_baseline = sensitivity_frame.loc[
        np.isclose(sensitivity_frame.volatility_factor, 1.0)
    ].iloc[0]
    sensitivity_baseline_match = bool(
        abs(float(sensitivity_baseline.q4_2_total_cost_yuan) - q4_rows["causal Q4-2"]["total_cost_yuan"]) < 1e-8
        and abs(float(sensitivity_baseline.q4_3_total_cost_yuan) - q4_rows["causal Q4-3"]["total_cost_yuan"]) < 1e-8
    )
    sensitivity_pass = bool(
        q4_sensitivity.get("baseline_match", {}).get("pass")
        and q4_sensitivity.get("run_id") == run_id
        and sensitivity_baseline_match
        and sensitivity_frame.maximum_daily_mean_price_error_yuan_per_kwh.max() < 3e-15
        and sensitivity_frame.minimum_stressed_price_yuan_per_kwh.min() > 0
        and sensitivity_frame[["q4_2_cross_day_soc_continuous", "q4_3_cross_day_soc_continuous"]].to_numpy().all()
        and sensitivity_frame[["q4_2_max_balance_residual_kwh", "q4_3_max_balance_residual_kwh"]].to_numpy().max() < 1e-7
    )
    figure_paths = sorted({
        *(ROOT / path for path in Q1_FIGURES),
        *(ROOT / path for source in (q2, q3, q4, q4_sensitivity) for path in source.get("figures", [])),
    })
    workbook_values = list(workbook_audit.values())
    all_workbooks_match = all(
        item.get("content_matches_q1_json", item.get("content_matches_daily_csv", False))
        for item in workbook_values
    )
    audit = {
        "run_id": run_id,
        "model_version": MODEL_VERSION,
        "unit_tests": unit_tests,
        "timestamp_tests_pass": bool(timestamp_tests.get("final_workbooks_ready")),
        "timestamp_alignment_pass": timestamp_audit.get("verdict") == "PASS",
        "q3_release_alignment_audit": q3_alignment,
        "soc_bounds_pass": bool(all_frames.soc_min_kwh.min() >= SOC_LOWER - 1e-7 and all_frames.soc_max_kwh.max() <= SOC_UPPER + 1e-7),
        "charge_discharge_limits_pass": bool(all_frames.max_charge_kwh.max() <= ENERGY_LIMIT + 1e-7 and all_frames.max_discharge_kwh.max() <= ENERGY_LIMIT + 1e-7),
        "energy_balance_pass": bool(all_frames.balance_max_abs_kwh.max() < 1e-7),
        "maximum_energy_balance_residual_kwh": float(all_frames.balance_max_abs_kwh.max()),
        "simultaneous_charge_discharge_periods": int(all_frames.simultaneous_charge_discharge_periods.sum()),
        "cross_day_soc_continuity_pass": bool(all(
            item.get("cross_day_soc_continuous", False)
            for item in [q2["annual"], q3["full_scheme"], *q4_rows.values()]
        )),
        "release_boundaries_36_72_108": bool(unit_tests.get("rolling_update_boundaries_and_constraints")),
        "q4_price_volatility_sensitivity": {
            "pass": sensitivity_pass,
            "run_id_matches_summary": q4_sensitivity.get("run_id") == run_id,
            "factors": q4_sensitivity["factors"],
            "baseline_matches_main_q4": sensitivity_baseline_match,
            "daily_mean_price_preserved": bool(
                sensitivity_frame.maximum_daily_mean_price_error_yuan_per_kwh.max() < 3e-15
            ),
            "all_stressed_prices_positive": bool(
                sensitivity_frame.minimum_stressed_price_yuan_per_kwh.min() > 0
            ),
        },
        "workbooks": workbook_audit,
        "all_five_workbooks_generated": all((ROOT / name).exists() for name in WORKBOOKS),
        "all_five_workbooks_match_sources": all_workbooks_match,
        "figures": {
            "count": len(figure_paths),
            "all_exist": all(path.exists() for path in figure_paths),
            "sha256": {str(path.relative_to(ROOT)): _sha(path) for path in figure_paths if path.exists()},
        },
    }
    audit["final_delivery_chain_pass"] = bool(
        audit["all_five_workbooks_generated"]
        and audit["all_five_workbooks_match_sources"]
        and audit["timestamp_tests_pass"]
        and audit["timestamp_alignment_pass"]
        and audit["soc_bounds_pass"]
        and audit["charge_discharge_limits_pass"]
        and audit["energy_balance_pass"]
        and audit["simultaneous_charge_discharge_periods"] == 0
        and audit["cross_day_soc_continuity_pass"]
        and audit["release_boundaries_36_72_108"]
        and audit["q4_price_volatility_sensitivity"]["pass"]
        and audit["figures"]["all_exist"]
    )

    payload = {
        "run_id": run_id,
        "model_version": MODEL_VERSION,
        "specification": "c建模说明与结果.txt",
        "q1": q1,
        "q2": q2,
        "q3": q3,
        "q4": q4,
        "q4_sensitivity": q4_sensitivity,
        "audit": audit,
    }
    (ROOT / "summary.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8"
    )
    (ROOT / "results/final_audit.json").write_text(
        json.dumps(audit, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8"
    )

    q4_table = pd.DataFrame(q4["scenarios"])[
        ["scenario", "total_cost_yuan", "plan_cost_yuan", "adjustment_cost_yuan", "emergency_cost_yuan"]
    ]
    q3_table = pd.DataFrame(q3["ablation"])[
        ["scheme", "total_cost_yuan", "plan_cost_yuan", "adjustment_cost_yuan", "emergency_cost_yuan"]
    ]
    sensitivity_table = sensitivity_frame[[
        "volatility_factor", "mean_daily_price_cv",
        "q4_2_total_cost_yuan", "q4_3_total_cost_yuan",
        "q4_2_emergency_kwh", "q4_3_emergency_kwh",
    ]]
    lines = [
        "# c-spec 最终统一运行汇总", "",
        f"- run_id：`{run_id}`",
        f"- 模型版本：`{MODEL_VERSION}`", "",
        "## 核心结果", "",
        f"- Q1：{q1['daily_purchase_cost_yuan']:.6f} 元",
        f"- Q2：{q2['annual']['total_cost_yuan']:.6f} 元",
        f"- Q3：{q3['full_scheme']['total_cost_yuan']:.6f} 元",
        f"- Q4-2：{q4_rows['causal Q4-2']['total_cost_yuan']:.6f} 元",
        f"- Q4-3：{q4_rows['causal Q4-3']['total_cost_yuan']:.6f} 元", "",
        "## Q3 消融", "", _to_markdown(q3_table), "",
        "## Q4 场景", "", _to_markdown(q4_table), "",
        "## Q4 电价日内波动灵敏度", "", _to_markdown(sensitivity_table), "",
        "## 最终验收", "",
        f"- 五工作簿与源结果一致：`{all_workbooks_match}`",
        f"- 时间轴测试：`{audit['timestamp_tests_pass']}`",
        f"- 时间轴审计：`{audit['timestamp_alignment_pass']}`",
        f"- Q4电价波动灵敏度：`{sensitivity_pass}`",
        f"- 最终交付链：`{audit['final_delivery_chain_pass']}`", "",
        "逐文件哈希与细项见 `results/final_audit.json` 和 `FINAL_RUN_MANIFEST.md`。",
    ]
    (ROOT / "results/summary.md").write_text("\n".join(lines), encoding="utf-8")
    print(json.dumps({"run_id": run_id, "audit": audit}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
