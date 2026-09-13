"""Regression gate for the accepted j-to-j Attachment 5 output mapping."""
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "src"), str(ROOT)]

import problem1  # noqa: E402
from config import MODEL_VERSION, N  # noqa: E402
from result_workbooks import OFFICIAL_SHEETS  # noqa: E402
from scenario_cache import cache_path  # noqa: E402
from time_axis import all_slot_labels, interval_start_index  # noqa: E402

FILES = ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx")
KEY_LABELS = {
    0: "0:00-0:10",
    35: "5:50-6:00",
    36: "6:00-6:10",
    60: "10:00-10:10",
    72: "12:00-12:10",
    108: "18:00-18:10",
    143: "23:50-0:00+1",
}
CACHE_SPECS = {
    "result2.xlsx": ("q2_fixed", {"model": MODEL_VERSION, "price": "attachment1"}, False),
    "result3.xlsx": (
        "q3_6_12_18_rec",
        {"model": MODEL_VERSION, "updates": [6, 12, 18], "recourse_plan": True},
        True,
    ),
    "result4-2.xlsx": ("q4_causal_q2", {"model": MODEL_VERSION, "price": "causal"}, False),
    "result4-3.xlsx": (
        "q4_causal_q3",
        {"model": MODEL_VERSION, "price": "causal", "updates": [6, 12, 18], "recourse": True},
        True,
    ),
}


def _slot_labels(sheet, vertical: bool = False) -> list[str]:
    if vertical:
        return [sheet.cell(position + 2, 1).value for position in range(N)]
    return [sheet.cell(1, position + 2).value for position in range(N)]


def _style_signature(cell) -> tuple:
    return (
        cell.style_id,
        cell.number_format,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.alignment.wrap_text,
    )


def _template_style_checks(filename: str, workbook) -> dict:
    template = openpyxl.load_workbook(ROOT / "data" / "附件" / "附件5" / filename)
    checks = {}
    for sheet_name in OFFICIAL_SHEETS[filename]:
        generated = workbook[sheet_name]
        source = template[sheet_name]
        assert generated.sheet_state == source.sheet_state
        assert list(generated.merged_cells.ranges) == list(source.merged_cells.ranges)
        for column in range(1, min(source.max_column, N + 3) + 1):
            assert generated.column_dimensions[openpyxl.utils.get_column_letter(column)].width == source.column_dimensions[openpyxl.utils.get_column_letter(column)].width
        rows_to_compare = range(1, source.max_row + 1)
        if filename != "result1.xlsx" and sheet_name in ("充放电量", "紧急购电量"):
            # These two sheets contain abbreviated example rows plus an ellipsis
            # in Attachment 5.  The generated workbook expands those examples.
            rows_to_compare = range(1, 2)
        for row in rows_to_compare:
            assert generated.row_dimensions[row].height == source.row_dimensions[row].height
            for column in range(1, source.max_column + 1):
                actual_style = _style_signature(generated.cell(row, column))
                expected_style = _style_signature(source.cell(row, column))
                assert actual_style == expected_style, (
                    f"style mismatch {filename}/{sheet_name}!{generated.cell(row, column).coordinate}: "
                    f"{actual_style} != {expected_style}"
                )

        # The annual storage/emergency sheets grow beyond the official example
        # rows.  Every added row must inherit the corresponding official style.
        if filename != "result1.xlsx" and sheet_name == "充放电量":
            for row in range(2, generated.max_row + 1):
                template_row = 2 + ((row - 2) % 6)
                assert generated.row_dimensions[row].height == source.row_dimensions[template_row].height
                for column in range(1, 7):
                    assert _style_signature(generated.cell(row, column)) == _style_signature(source.cell(template_row, column))
        if filename != "result1.xlsx" and sheet_name == "紧急购电量":
            for row in range(2, generated.max_row + 1):
                assert generated.row_dimensions[row].height == source.row_dimensions[2].height
                for column in range(1, 4):
                    assert _style_signature(generated.cell(row, column)) == _style_signature(source.cell(2, column))
        checks[sheet_name] = True
    template.close()
    return checks


def _verify_q1(workbook) -> dict:
    raw, price, load, pv = problem1.load_inputs()
    solution = problem1.solve_lp(price, load, pv)
    plan = workbook["计划购电量"]
    labels = _slot_labels(plan, vertical=True)
    values = np.asarray([plan.cell(position + 2, 2).value for position in range(N)], dtype=float)
    np.testing.assert_allclose(values, solution["purchase"], rtol=0, atol=1e-9)
    storage = workbook["充放电量"]
    charge = np.asarray([storage.cell(row, 2).value for row in range(2, 8)], dtype=float)
    discharge = np.asarray([storage.cell(row, 3).value for row in range(2, 8)], dtype=float)
    expected_charge = np.asarray([solution["charge"][block * 24:(block + 1) * 24].sum() for block in range(6)])
    expected_discharge = np.asarray([solution["discharge"][block * 24:(block + 1) * 24].sum() for block in range(6)])
    np.testing.assert_allclose(charge, expected_charge, rtol=0, atol=1e-9)
    np.testing.assert_allclose(discharge, expected_discharge, rtol=0, atol=1e-9)
    assert abs(float(storage.cell(2, 5).value) - float(solution["soc"][0])) < 1e-9
    assert abs(float(storage.cell(3, 5).value) - float(solution["soc"][-1])) < 1e-9
    return {
        "labels": {str(position): labels[position] for position in KEY_LABELS},
        "position_j_equals_internal_j": True,
        "purchase_max_abs_error_kwh": float(np.max(np.abs(values - solution["purchase"]))),
        "cost_yuan": float(solution["cost"]),
        "source_first_timestamp": str(raw.iloc[0, 0]),
    }


def _verify_annual(filename: str, workbook, spec: tuple) -> dict:
    label, options, rolling = spec
    path = cache_path(label, options)
    assert path.exists(), f"missing current-run cache: {path}"
    plan_sheet = workbook["计划购电量"]
    labels = _slot_labels(plan_sheet)
    with np.load(path, allow_pickle=False) as data:
        dates = pd.DatetimeIndex(pd.to_datetime(data["date"]))
        first = int(np.flatnonzero(dates == pd.Timestamp("2025-02-01"))[0])
        ids = np.arange(first, len(dates))
        workbook_dates = pd.DatetimeIndex(
            pd.to_datetime([plan_sheet.cell(row, 1).value for row in range(2, 2 + len(ids))])
        )
        assert workbook_dates.equals(dates[ids])
        plan = np.asarray([
            [plan_sheet.cell(row, position + 2).value for position in range(N)]
            for row in range(2, 2 + len(ids))
        ], dtype=float)
        expected_plan = data["plan"][ids]
        np.testing.assert_allclose(plan, expected_plan, rtol=0, atol=1e-9)
        final_error = 0.0
        release_boundaries_pass = not rolling
        if rolling:
            adjustment = workbook["调整购电量"]
            assert _slot_labels(adjustment) == labels
            final = np.asarray([
                [adjustment.cell(row, position + 2).value for position in range(N)]
                for row in range(2, 2 + len(ids))
            ], dtype=float)
            expected_final = data["final"][ids]
            np.testing.assert_allclose(final, expected_final, rtol=0, atol=1e-9)
            final_error = float(np.max(np.abs(final - expected_final)))
            hours = data["update_hours"][ids]
            assert np.all(hours == np.asarray([6, 12, 18]))
            snapshots = data["update_purchase"][ids]
            for update_index, boundary in enumerate((36, 72, 108)):
                assert np.isnan(snapshots[:, update_index, :boundary]).all()
                assert np.isfinite(snapshots[:, update_index, boundary:]).all()
            release_boundaries_pass = True
    return {
        "cache": str(path.relative_to(ROOT)),
        "labels": {str(position): labels[position] for position in KEY_LABELS},
        "position_j_equals_internal_j": True,
        "plan_max_abs_error_kwh": float(np.max(np.abs(plan - expected_plan))),
        "final_max_abs_error_kwh": final_error,
        "release_boundaries_36_72_108": release_boundaries_pass,
    }


def audit_workbooks(workbook_root: Path) -> dict:
    expected_labels = all_slot_labels()
    assert len(expected_labels) == N
    for position, label in KEY_LABELS.items():
        assert expected_labels[position] == label
    assert [interval_start_index(hour) for hour in (6, 12, 18)] == [36, 72, 108]

    results = {}
    for filename in FILES:
        path = workbook_root / filename
        assert path.exists(), f"missing workbook: {path}"
        workbook = openpyxl.load_workbook(path, read_only=False, data_only=True)
        assert workbook.sheetnames == OFFICIAL_SHEETS[filename]
        plan = workbook["计划购电量"]
        labels = _slot_labels(plan, vertical=filename == "result1.xlsx")
        assert labels == expected_labels
        style_checks = _template_style_checks(filename, workbook)
        detail = _verify_q1(workbook) if filename == "result1.xlsx" else _verify_annual(filename, workbook, CACHE_SPECS[filename])
        detail.update({
            "official_sheet_set": True,
            "template_styles_preserved": all(style_checks.values()),
            "all_144_labels_exact": True,
        })
        results[filename] = detail
        workbook.close()
    return {
        "mapping_rule": "template position j <- internal j; labels changed, arrays unchanged",
        "key_labels": {str(position): label for position, label in KEY_LABELS.items()},
        "release_boundaries": {"06:00": 36, "12:00": 72, "18:00": 108},
        "workbooks": results,
        "final_workbooks_ready": True,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook-root", type=Path, default=ROOT)
    parser.add_argument("--output", type=Path, default=ROOT / "results" / "timestamp_alignment_tests.json")
    args = parser.parse_args()
    result = audit_workbooks(args.workbook_root.resolve())
    output = args.output if args.output.is_absolute() else ROOT / args.output
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
