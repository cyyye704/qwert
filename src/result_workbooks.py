"""Official-template Excel exports for the current c-spec results.

The exporter changes display labels and writes values position-for-position. It
never shifts, rolls, slices with an offset, or borrows a value from another day.
"""
from __future__ import annotations

from copy import copy
from pathlib import Path

import numpy as np
import openpyxl

from config import N
from time_axis import all_slot_labels, interval_label

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "data" / "附件" / "附件5"

OFFICIAL_SHEETS = {
    "result1.xlsx": ["计划购电量", "充放电量"],
    "result2.xlsx": ["计划购电量", "充放电量", "紧急购电量"],
    "result3.xlsx": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
    "result4-2.xlsx": ["计划购电量", "充放电量", "紧急购电量"],
    "result4-3.xlsx": ["计划购电量", "调整购电量", "充放电量", "紧急购电量"],
}


def _output_path(filename: str, output_path: str | Path | None) -> Path:
    output = ROOT / filename if output_path is None else Path(output_path)
    if not output.is_absolute():
        output = ROOT / output
    output.parent.mkdir(parents=True, exist_ok=True)
    return output


def _load_official_template(filename: str):
    if filename not in OFFICIAL_SHEETS:
        raise ValueError(f"unsupported official workbook: {filename}")
    workbook = openpyxl.load_workbook(TEMPLATES / filename)
    if workbook.sheetnames != OFFICIAL_SHEETS[filename]:
        raise AssertionError(
            f"official sheet set changed for {filename}: {workbook.sheetnames}"
        )
    return workbook


def _copy_row_style(ws, source_row: int, target_row: int, max_column: int) -> None:
    """Copy only presentation metadata from one official-template row."""
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    ws.row_dimensions[target_row].hidden = ws.row_dimensions[source_row].hidden
    for column in range(1, max_column + 1):
        source = ws.cell(source_row, column)
        target = ws.cell(target_row, column)
        if source.has_style:
            target._style = copy(source._style)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)


def _set_horizontal_slot_labels(ws, start_column: int = 2) -> None:
    labels = all_slot_labels()
    if len(labels) != N:
        raise AssertionError(f"expected {N} labels, got {len(labels)}")
    for position, label in enumerate(labels):
        ws.cell(1, start_column + position).value = label


def _finite_vector(value, name: str) -> np.ndarray:
    array = np.asarray(value, dtype=float)
    if array.shape != (N,) or not np.isfinite(array).all():
        raise AssertionError(f"{name} must be a finite ({N},) vector")
    return array


def _emergency_rows(record):
    values = _finite_vector(record["actual"]["emergency"], "emergency")
    mask = values > 1e-8
    starts = np.flatnonzero(mask & ~np.r_[False, mask[:-1]])
    ends = np.flatnonzero(mask & ~np.r_[mask[1:], False]) + 1
    return [
        [record["date"], interval_label(int(start), int(end)), float(values[start:end].sum())]
        for start, end in zip(starts, ends)
    ]


def export_problem1(solution: dict, output_path: str | Path | None = None) -> Path:
    """Write the current Q1 solution into a fresh official result1 template."""
    workbook = _load_official_template("result1.xlsx")
    purchase = _finite_vector(solution["purchase"], "Q1 purchase")
    charge = _finite_vector(solution["charge"], "Q1 charge")
    discharge = _finite_vector(solution["discharge"], "Q1 discharge")
    soc = np.asarray(solution["soc"], dtype=float)
    if soc.shape != (N + 1,) or not np.isfinite(soc).all():
        raise AssertionError(f"Q1 soc must be a finite ({N + 1},) vector")

    plan = workbook["计划购电量"]
    labels = all_slot_labels()
    for position, (label, value) in enumerate(zip(labels, purchase)):
        row = position + 2
        plan.cell(row, 1).value = label
        plan.cell(row, 2).value = float(value)

    storage = workbook["充放电量"]
    for block in range(6):
        start, end = block * N // 6, (block + 1) * N // 6
        row = block + 2
        storage.cell(row, 2).value = float(charge[start:end].sum())
        storage.cell(row, 3).value = float(discharge[start:end].sum())
    storage.cell(2, 5).value = float(soc[0])
    storage.cell(3, 5).value = float(soc[-1])

    output = _output_path("result1.xlsx", output_path)
    workbook.save(output)
    return output


def export_records(
    records: list[dict],
    metrics,
    filename: str,
    rolling: bool = False,
    output_path: str | Path | None = None,
) -> Path:
    """Write Q2/Q3/Q4 records into a fresh official template.

    For rolling workbooks, the official one-row-per-day ``调整购电量`` sheet
    contains the final piecewise purchase schedule after the 06/12/18 updates.
    The three release snapshots remain available in JSON/cache audit artifacts.
    """
    workbook = _load_official_template(filename)
    if len(records) != len(metrics):
        raise AssertionError(f"record/metric length mismatch: {len(records)} != {len(metrics)}")
    if rolling != ("调整购电量" in workbook.sheetnames):
        raise AssertionError(f"rolling/template mismatch for {filename}")

    plan_sheet = workbook["计划购电量"]
    _set_horizontal_slot_labels(plan_sheet)
    adjustment_sheet = workbook["调整购电量"] if rolling else None
    if adjustment_sheet is not None:
        _set_horizontal_slot_labels(adjustment_sheet)

    storage_sheet = workbook["充放电量"]
    emergency_sheet = workbook["紧急购电量"]
    storage_labels = [storage_sheet.cell(2 + block, 2).value for block in range(6)]
    storage_rows: list[list] = []
    emergency_rows: list[list] = []

    metric_rows = list(metrics.iterrows())
    for row_index, (record, (_, metric)) in enumerate(zip(records, metric_rows), start=2):
        base = _finite_vector(record["plan"]["purchase"], "plan purchase")
        plan_sheet.cell(row_index, 1).value = record["date"]
        for position, value in enumerate(base):
            plan_sheet.cell(row_index, position + 2).value = float(value)
        plan_sheet.cell(row_index, N + 2).value = float(base.sum())
        plan_sheet.cell(row_index, N + 3).value = float(metric.plan_cost_yuan)

        actual = record["actual"]
        charge = _finite_vector(actual["charge"], "actual charge")
        discharge = _finite_vector(actual["discharge"], "actual discharge")
        soc = np.asarray(actual["soc"], dtype=float)
        if soc.shape != (N + 1,) or not np.isfinite(soc).all():
            raise AssertionError(f"actual soc must be a finite ({N + 1},) vector")
        for block in range(6):
            start, end = block * N // 6, (block + 1) * N // 6
            storage_rows.append([
                record["date"] if block == 0 else None,
                storage_labels[block],
                float(charge[start:end].sum()),
                float(discharge[start:end].sum()),
                "0:00" if block == 0 else ("24:00" if block == 1 else None),
                float(soc[0]) if block == 0 else (float(soc[-1]) if block == 1 else None),
            ])
        emergency_rows.extend(_emergency_rows(record))

        if adjustment_sheet is not None:
            final_purchase = _finite_vector(record["final_purchase"], "final purchase")
            adjustment_sheet.cell(row_index, 1).value = record["date"]
            for position, value in enumerate(final_purchase):
                adjustment_sheet.cell(row_index, position + 2).value = float(value)
            adjustment_sheet.cell(row_index, N + 2).value = float(final_purchase.sum())
            adjustment_sheet.cell(row_index, N + 3).value = float(
                metric.plan_cost_yuan + metric.adjustment_cost_yuan
            )

    for target_row, values in enumerate(storage_rows, start=2):
        source_row = 2 + ((target_row - 2) % 6)
        _copy_row_style(storage_sheet, source_row, target_row, 6)
        for column, value in enumerate(values, start=1):
            storage_sheet.cell(target_row, column).value = value

    for target_row, values in enumerate(emergency_rows, start=2):
        _copy_row_style(emergency_sheet, 2, target_row, 3)
        for column, value in enumerate(values, start=1):
            emergency_sheet.cell(target_row, column).value = value

    output = _output_path(filename, output_path)
    workbook.save(output)
    return output
