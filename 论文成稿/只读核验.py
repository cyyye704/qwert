"""独立复算原始结果的约束和费用；不调用优化器，不修改任何原文件。"""
from pathlib import Path
import hashlib
import json
import re
import sys
import zipfile
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pypdf import PdfReader

sys.stdout.reconfigure(encoding='utf-8')
OUT = Path(__file__).resolve().parent
ROOT = OUT.parent
report = {}
base = json.loads((OUT/'原文件校验基线.json').read_text(encoding='utf-8'))
changed = [name for name, h in base.items() if hashlib.sha256((ROOT/name).read_bytes()).hexdigest() != h]
assert not changed, changed
report['original_files_unchanged'] = len(base)

raw1 = pd.read_excel(ROOT/'data/附件/附件1.xlsx')
load = pd.read_excel(ROOT/'data/附件/附件2.xlsx', sheet_name='小区负载').iloc[:, 1:].to_numpy(float)/6
pv = pd.read_excel(ROOT/'data/附件/附件2.xlsx', sheet_name='光伏发电实际功率').iloc[:, 1:].to_numpy(float)/6
price = pd.read_excel(ROOT/'data/附件/附件4.xlsx').iloc[:, 1:].to_numpy(float)
fixed = raw1.iloc[:, 1].to_numpy(float)
q1 = json.loads((ROOT/'results/problem1.json').read_text(encoding='utf-8'))
w = load_workbook(ROOT/'result1.xlsx', read_only=True, data_only=True)
g1 = np.array([r[1] for r in list(w['计划购电量'].values)[1:]], float)
w.close()
assert abs(g1.sum()-q1['daily_purchase_kwh']) < 1e-7
assert abs(g1@fixed-q1['daily_purchase_cost_yuan']) < 1e-7
report['q1_cost_recomputed_yuan'] = float(g1@fixed)

cache_results = []
cache_data = {}
for path in sorted((ROOT/'build/c_spec_cache').glob('*.npz')):
    with np.load(path, allow_pickle=False) as z:
        a = {key: z[key] for key in z.files}
    assert a['soc'].shape == (365, 145)
    c, d, s = a['charge'], a['discharge'], a['soc']
    bal = a['final']+pv+d+a['emergency']-load-c-a['waste']
    rec = np.diff(s, axis=1)-.9*c+d/.9
    assert np.max(np.abs(bal)) < 1e-7
    assert np.max(np.abs(rec)) < 1e-7
    assert np.max(np.abs(s[:-1, -1]-s[1:, 0])) < 1e-7
    assert s.min() >= 1200-1e-7 and s.max() <= 10800+1e-7
    assert c.min() >= -1e-8 and d.min() >= -1e-8
    assert c.max() <= 5000/6+1e-7 and d.max() <= 5000/6+1e-7
    assert not np.any((c>1e-8) & (d>1e-8))
    assert abs(s[0,0]-6000) < 1e-7
    cache_results.append({'file': path.name, 'max_balance_recomputed_kwh': float(np.max(np.abs(bal))), 'max_soc_recursion_kwh': float(np.max(np.abs(rec))), 'cross_day_continuous': True})
    cache_data[path.name] = a
report['all_cache_audits'] = cache_results

rows4 = {r['scenario']: r for r in json.loads((ROOT/'results/problem4.json').read_text(encoding='utf-8'))['scenarios']}
models = [
    ('2', 'q2_fixed_', 'result2.xlsx', 'problem2_daily.csv', False, fixed),
    ('3', 'q3_6_12_18_rec_', 'result3.xlsx', 'problem3_daily.csv', True, fixed),
    ('42', 'q4_causal_q2_', 'result4-2.xlsx', 'problem4_causal_q2_daily.csv', False, price),
    ('43', 'q4_causal_q3_', 'result4-3.xlsx', 'problem4_causal_q3_daily.csv', True, price),
]
annual = []
for mode, prefix, workbook, daily, rolling, settlement in models:
    name = next(name for name in cache_data if name.startswith(prefix))
    a = cache_data[name]
    ids = np.flatnonzero(a['date'] >= '2025-02-01')
    assert len(ids) == 334
    p = np.tile(settlement, (365,1)) if settlement.ndim == 1 else settlement
    cp = np.sum(a['plan']*p, axis=1)
    ce = np.sum(5*a['emergency']*p, axis=1)
    ca = np.zeros(365)
    if rolling:
        ca = np.sum(p*(-.5*np.maximum(a['plan']-a['final'],0)+1.5*np.maximum(a['final']-a['plan'],0)), axis=1)
        assembled = a['plan'].copy()
        for j,h in enumerate([6,12,18]):
            assert np.all(a['update_hours'][:,j] == h)
            assembled[:, h*6:] = a['update_purchase'][:,j,h*6:]
        assert np.max(np.abs(assembled-a['final'])) < 1e-9
    df = pd.read_csv(ROOT/'results'/daily)
    for key, v in [('plan_cost_yuan', cp), ('adjustment_cost_yuan', ca), ('emergency_cost_yuan', ce), ('total_cost_yuan', cp+ca+ce)]:
        assert np.max(np.abs(v[ids]-df[key].to_numpy(float))) < 1e-6
    w = load_workbook(ROOT/workbook, read_only=True, data_only=True)
    vals = list(w['计划购电量'].values)[1:]
    dates = [str(r[0])[:10] for r in vals]
    assert dates == a['date'][ids].tolist()
    wg = np.array([r[1:145] for r in vals], float)
    err = float(np.max(np.abs(wg-a['plan'][ids])))
    assert err < 1e-8
    if rolling:
        vals = list(w['调整购电量'].values)[1:]
        wf = np.array([r[1:145] for r in vals], float)
        assert np.max(np.abs(wf-a['final'][ids])) < 1e-8
    w.close()
    q = json.loads((ROOT/f'results/problem{mode[0]}.json').read_text(encoding='utf-8'))
    typical = q['typical_days'] if len(mode)==1 else q['typical_days'][f'q4_{mode[-1]}']
    for date, t in typical.items():
        i = list(a['date']).index(date)
        for key, ck in [('plan_purchase_kwh','plan'),('final_purchase_kwh','final'),('charge_kwh','charge'),('discharge_kwh','discharge'),('soc_kwh','soc'),('emergency_kwh','emergency')]:
            assert np.max(np.abs(np.array(t[key])-a[ck][i])) < 1e-8
        assert abs(t['metrics']['total_cost_yuan']-(cp+ca+ce)[i]) < 1e-6
    annual.append({'model': mode, 'cost_yuan': float((cp+ca+ce)[ids].sum()), 'emergency_kwh': float(a['emergency'][ids].sum()), 'workbook_cache_max_difference_kwh': err, 'typical_arrays_match_cache': True})
report['annual_recomputation'] = annual

sens = pd.read_csv(ROOT/'results/problem1_sensitivity.csv', index_col=0).to_numpy(float)
assert sens.shape == (11,11)
assert (np.diff(sens, axis=0) <= 1e-6).all() and (np.diff(sens,axis=1) <= 1e-6).all()
report['q1_sensitivity_cost_range'] = [float(sens.min()),float(sens.max())]
vr = pd.read_csv(ROOT/'results/problem4_price_volatility_sensitivity.csv')
assert (vr['q4_3_total_cost_yuan'] < vr['q4_2_total_cost_yuan']).all()
report['q4_volatility_daily_mean_max_error'] = float(vr['maximum_daily_mean_price_error_yuan_per_kwh'].max())

md = (OUT/'微网电力调控策略_论文.md').read_text(encoding='utf-8')
assert '@@' not in md and '\x08' not in md
assert [int(x) for x in re.findall(r'^表(\d+) ', md, re.M)] == list(range(1,38))
assert [int(x) for x in re.findall(r'\\tag\{(\d+)\}',md)] == list(range(1,21))
assert set(re.findall(r'^\[(\d+)\]',md,re.M)) == set('123456')
body = md.split('## 参考文献')[0]
for i in range(1,7):
    assert f'[{i}]' in body
for v in ['35101.57','59482.70','13701658.50','13266622.73','14401772.58','13953460.93','13155783.13']:
    assert v in md
abstract = md.split('## 摘要')[1].split('**关键词')[0]
report['abstract_characters_without_whitespace'] = len(re.sub(r'\s','',abstract))

pdf = PdfReader(OUT/'微网电力调控策略_论文.pdf')
report['pdf_pages'] = len(pdf.pages)
assert len(pdf.pages) >= 20
pdftext = '\n'.join(page.extract_text() for page in pdf.pages)
assert '问题四' in pdftext and '附录' in pdftext
report['pdf_all_pages_have_content'] = all(len(page.extract_text())>200 for page in pdf.pages)
log = (OUT/'微网电力调控策略_论文.log').read_text(encoding='utf-8',errors='replace')
report['overfull_boxes'] = log.count('Overfull')
report['missing_glyphs'] = log.count('Missing character')
assert report['overfull_boxes'] == 0 and report['missing_glyphs'] == 0

with zipfile.ZipFile(OUT/'原始程序与结果.zip') as z:
    assert z.testzip() is None
    for name in z.namelist():
        if (ROOT/name).is_file():
            assert z.read(name) == (ROOT/name).read_bytes()
    report['source_package_entries'] = len(z.namelist())
report['solver_rerun'] = False
(OUT/'只读核验结果.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='all_cache_audits'},ensure_ascii=False,indent=2))
